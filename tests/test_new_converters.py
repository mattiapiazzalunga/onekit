# Copyright 2026 Mattia Piazzalunga
# SPDX-License-Identifier: LicenseRef-PolyForm-Shield-1.0.0
import csv
import io
import json
import shutil
import subprocess
import zipfile

import pytest
from openpyxl import Workbook, load_workbook
from PIL import Image
from pypdf import PdfReader

from onekit import Artifact, OneKitError, execute, pipeline

IDS = ['xml-to-xlsx','xlsx-to-xml','xml-to-csv','csv-to-xml','xml-to-json','json-to-csv',
       'json-to-xml','yaml-to-json','json-to-yaml','tiff-to-svg','epub-to-pdf',
       'pdf-tables-to-xlsx','pdf-tables-to-csv','svg-to-png']


def samples(tool):
    if tool.id.startswith('pdf-tables-to-'):
        from reportlab.pdfgen.canvas import Canvas
        stream = io.BytesIO();canvas = Canvas(stream)
        for i,row in enumerate([['Name','Count','Price'],['Apples','2','5'],['Pears','3','7']]):
            for x,text in zip([50,220,370],row):
                canvas.drawString(x,700-24*i,text)
        canvas.save()
        return [Artifact('tables.pdf',stream.getvalue())],{}
    source = tool.id.split('-to-')[0]
    if source == 'xml':
        return [Artifact('data.xml',b'<rows><row><name>Ada</name><count>2</count></row><row><name>Grace</name><count>3</count></row></rows>')],{}
    if source == 'json':
        return [Artifact('data.json',b'[{"name":"Ada","count":2},{"name":"Grace","count":3}]')],{}
    if source == 'yaml':
        return [Artifact('data.yaml',b'name: Ada\ncount: 2\nactive: true\n')],{}
    if source == 'csv':
        return [Artifact('data.csv',b'name,count\nAda,2\nGrace,3\n')],{}
    if source == 'xlsx':
        book = Workbook();book.active.append(['name','count']);book.active.append(['Ada',2])
        stream = io.BytesIO();book.save(stream)
        return [Artifact('data.xlsx',stream.getvalue())],{}
    if source == 'svg':
        return [Artifact('shape.svg',b'<svg xmlns="http://www.w3.org/2000/svg" width="64" height="48"><rect width="64" height="48" fill="#ff0000"/></svg>')],{}
    if source == 'tiff':
        stream = io.BytesIO();Image.new('RGB',(32,24),'blue').save(stream,'TIFF')
        return [Artifact('image.tiff',stream.getvalue())],{}
    stream = io.BytesIO()
    with zipfile.ZipFile(stream,'w') as archive:
        archive.writestr('mimetype','application/epub+zip')
        archive.writestr('META-INF/container.xml','<container><rootfiles><rootfile full-path="book/package.opf"/></rootfiles></container>')
        archive.writestr('book/package.opf','<package><manifest><item id="second" href="two.xhtml" media-type="application/xhtml+xml"/><item id="first" href="one.xhtml" media-type="application/xhtml+xml"/></manifest><spine><itemref idref="first"/><itemref idref="second"/></spine></package>')
        archive.writestr('book/one.xhtml','<html><body><h1>First chapter</h1><p>Hello OneKit.</p></body></html>')
        archive.writestr('book/two.xhtml','<html><body><h1>Second chapter</h1><p>The end.</p></body></html>')
    return [Artifact('book.epub',stream.getvalue())],{}


@pytest.mark.parametrize('identifier',IDS)
def test_new_converter_produces_expected_content(identifier):
    from onekit.catalog import TOOLS
    files,options = samples(TOOLS[identifier]);result = execute(identifier,files,options)
    assert result.files and all(file.data for file in result.files)
    if identifier == 'epub-to-pdf':
        text = '\n'.join(page.extract_text() for page in PdfReader(io.BytesIO(result.files[0].data)).pages)
        assert text.index('First chapter')<text.index('Second chapter')
        assert 'Hello OneKit.' in text
    if identifier.startswith('pdf-tables-to-'):
        assert result.details['tables'] == 1
        if result.files[0].kind == 'csv':
            rows = list(csv.reader(io.StringIO(result.files[0].text())))
        else:
            rows = list(load_workbook(io.BytesIO(result.files[0].data)).active.values)
        assert list(rows[1]) == ['Apples','2','5']
    if identifier == 'svg-to-png':
        image = Image.open(io.BytesIO(result.files[0].data)).convert('RGB')
        assert image.size == (64,48) and image.getpixel((30,20)) == (255,0,0)
    if identifier == 'tiff-to-svg':
        assert '<path' in result.files[0].text() and '<image' not in result.files[0].text()


def test_json_xml_roundtrip_preserves_arbitrary_keys_scalars_arrays_and_null():
    value = {'invalid xml key': ['Caffè',2,False,None,{'a&b':3.5}],'empty':{},'blank':[]}
    result = pipeline([{'tool':'json-to-xml'},{'tool':'xml-to-json'}],[Artifact('data.json',json.dumps(value).encode())])
    assert json.loads(result.files[0].text()) == value


def test_svg_scaling_changes_content_size_without_clipping_or_blank_padding():
    from onekit.catalog import TOOLS
    files,_ = samples(TOOLS['svg-to-png'])
    result = execute('svg-to-png',files,{'scale':2})
    image = Image.open(io.BytesIO(result.files[0].data)).convert('RGBA')
    assert image.size == (128,96)
    assert image.getpixel((100,80)) == (255,0,0,255)


def test_json_xml_rejects_characters_xml_cannot_represent():
    with pytest.raises(OneKitError,match='control characters'):
        execute('json-to-xml',[Artifact('data.json',b'{"a":"\\u0000"}')])


def test_xml_spreadsheet_roundtrip_preserves_headers_and_formula_strings():
    result = pipeline([{'tool':'csv-to-xml'},{'tool':'xml-to-xlsx'}],[Artifact('data.csv',b'name,note\nAda,=1+1\n')])
    cell = load_workbook(io.BytesIO(result.files[0].data)).active['B2']
    assert cell.value == '=1+1' and cell.data_type == 's'


@pytest.mark.parametrize('text',['a: 1\na: 2','a: .nan','a: &a [*a]','!!python/object/apply:os.system [echo bad]'])
def test_yaml_rejects_duplicate_nonfinite_cyclic_or_unsafe_values(text):
    with pytest.raises(OneKitError):
        execute('yaml-to-json',[Artifact('data.yaml',text.encode())])


def test_xml_rejects_entities_and_reports_nested_table_shapes():
    with pytest.raises(OneKitError,match='entity'):
        execute('xml-to-json',[Artifact('data.xml',b'<!DOCTYPE a [<!ENTITY x "expanded">]><a>&x;</a>')])
    with pytest.raises(OneKitError,match='flat'):
        execute('xml-to-csv',[Artifact('data.xml',b'<rows><row><a><b>1</b></a></row></rows>')])


@pytest.mark.parametrize('content',['<image href="https://example.com/a.png"/>','<script>alert(1)</script>',
                                  '<rect style="fill:url(https://example.com/a.svg)"/>'])
def test_svg_external_content_is_rejected_before_renderer(content):
    with pytest.raises(OneKitError):
        execute('svg-to-png',[Artifact('bad.svg',f'<svg xmlns="http://www.w3.org/2000/svg" width="10" height="10">{content}</svg>'.encode())])


@pytest.mark.skipif(not shutil.which('ffmpeg'),reason='FFmpeg required')
def test_video_frames_are_multiple_decodable_images(tmp_path):
    source = tmp_path/'sample.mp4'
    subprocess.run(['ffmpeg','-v','error','-f','lavfi','-i','testsrc2=size=96x64:rate=10:duration=2','-c:v','libx264',str(source)],check=True)
    result = execute('video-to-frames',[Artifact('sample.mp4',source.read_bytes())],{'duration':1,'fps':2,'width':96})
    assert len(result.files) == 2
    assert all(Image.open(io.BytesIO(file.data)).size==(96,64) for file in result.files)
    all_frames = execute('video-to-frames',[Artifact('sample.mp4',source.read_bytes())],
                         {'duration':2,'all_frames':True,'width':96})
    assert len(all_frames.files) == 20
    variable = tmp_path/'variable.mp4'
    subprocess.run(['ffmpeg','-v','error','-i',str(source),'-vf',"select='not(mod(n,3))'",
                    '-fps_mode','vfr','-c:v','libx264',str(variable)],check=True)
    assert len(execute('video-to-frames',[Artifact('variable.mp4',variable.read_bytes())],
                       {'duration':2,'all_frames':True,'width':96}).files) == 7
    long_source = tmp_path/'long.mp4'
    subprocess.run(['ffmpeg','-v','error','-f','lavfi','-i','testsrc2=size=96x64:rate=10:duration=13',
                    '-c:v','libx264',str(long_source)],check=True)
    with pytest.raises(OneKitError,match='120 frames'):
        execute('video-to-frames',[Artifact('long.mp4',long_source.read_bytes())],
                {'duration':13,'all_frames':True,'width':96})


def test_source_frame_plan_and_passthrough_arguments():
    from onekit.tools.video_edit import arguments, source_frame_plan
    plan = source_frame_plan({'streams':[{'codec_type':'video','duration':'2',
                                            'avg_frame_rate':'30000/1001','nb_frames':'60'}]})
    assert plan['fps'] == pytest.approx(29.97002997)
    assert plan['frames'] == 60 and not plan['estimated']
    fallback = source_frame_plan({'format':{'duration':'2'},'streams':[{'codec_type':'video',
        'avg_frame_rate':'0/0','r_frame_rate':'60/1','nb_frames':'N/A'}]})
    assert fallback['frames'] == 120 and fallback['estimated']
    args = arguments('video-to-frames',{'all_frames':'true'})
    assert args[args.index('-vf')+1] == 'scale=1280:-1'
    assert args[args.index('-fps_mode')+1] == 'passthrough'
    assert args[args.index('-frames:v')+1] == '121'
    with pytest.raises(OneKitError):
        source_frame_plan({'streams':[]})
