# Copyright 2026 Mattia Piazzalunga
# SPDX-License-Identifier: LicenseRef-PolyForm-Shield-1.0.0
"""CLI adapter and output regressions."""
import io
import json

import pytest
from test_core import pdf

from onekit import Artifact, execute, pipeline
from onekit.catalog import TOOLS
from onekit.cli import doctor, main


def test_native_requirements_do_not_block_pure_python_converters(monkeypatch):
    def unexpected(name):
        pytest.fail(f"Pure converter should not require {name}")
    monkeypatch.setattr("onekit.tools.native.executable", unexpected)
    for id in ["md-to-json", "md-to-html", "html-to-json", "html-to-md", "xlsx-to-csv", "xlsx-to-json"]:
        assert not TOOLS[id].engines
        assert doctor(id)[0]["ready"]


def test_pdf_powerpoint_preserves_page_aspect_and_requested_resolution():
    from PIL import Image
    from pptx import Presentation
    from pypdf import PdfReader
    file = pdf()
    page_height = float(PdfReader(io.BytesIO(file.data)).pages[0].mediabox.height)
    for dpi in [36, 72]:
        result = execute("pdf-to-pptx", [file], {"dpi": dpi})
        presentation = Presentation(io.BytesIO(result.files[0].data))
        assert len(presentation.slides) == 2
        picture = presentation.slides[0].shapes[0]
        with Image.open(io.BytesIO(picture.image.blob)) as image:
            width, height = image.size
        assert height == round(page_height * dpi / 72)
        assert picture.width / picture.height == pytest.approx(width / height, rel=1e-5)
        assert picture.left == (presentation.slide_width-picture.width)//2
        assert picture.top == (presentation.slide_height-picture.height)//2
        assert result.details == {"pages": 2, "dpi": dpi}




@pytest.mark.native
@pytest.mark.parametrize("language,text,extension", [
    ("typescript", "const value:number=1;", "ts"),
    ("jsx", "const view=<div>Hello</div>;", "jsx"),
    ("tsx", "const view:unknown=<div>Hello</div>;", "tsx"),
])
def test_compiled_minifiers_output_javascript_and_work_in_cli_workflows(language, text, extension, tmp_path):
    source = Artifact(f"input.{extension}", text.encode())
    result = execute(f"dev-minify-{language}", [source])
    assert result.files[0].name == "input.js"
    assert TOOLS[f"dev-minify-{language}"].produces == "js"
    traced = pipeline([{"tool": f"dev-minify-{language}"}, {"tool": "dev-validate-javascript"}], [source])
    assert json.loads(traced.files[0].text())["valid"]
    path = tmp_path / source.name
    path.write_bytes(source.data)
    destination = tmp_path / "outputs"
    assert main(["run", f"dev-minify-{language}", str(path), "-o", str(destination)]) == 0
    assert (destination / "input.js").is_file()


@pytest.mark.native
def test_html_minifier_preserves_meaningful_inline_spaces_and_literal_blocks():
    text = '<p>Hello <span>world</span> !</p> <!-- omit --><pre>  a\n b</pre><script>globalThis.__executed=1;</script>'
    result = execute("dev-minify-html", [Artifact("input.html", text.encode())])
    from html.parser import HTMLParser
    class Text(HTMLParser):
        def __init__(self):
            super().__init__()
            self.values = []
        def handle_data(self, value):
            self.values.append(value)
    parser = Text()
    parser.feed(result.files[0].text())
    assert "".join(parser.values).startswith("Hello world !")
    assert "<pre>  a\n b</pre>" in result.files[0].text()
    assert "omit" not in result.files[0].text()


@pytest.mark.native
def test_formatters_preserve_selected_source_names():
    result = execute("dev-format-javascript", [Artifact("source.txt", b"const value=1;")])
    assert result.files[0].name == "source.txt"
