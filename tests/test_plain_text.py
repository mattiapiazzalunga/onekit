# Copyright 2026 Mattia Piazzalunga
# SPDX-License-Identifier: LicenseRef-PolyForm-Shield-1.0.0
import io

import pytest
from docx import Document

from onekit import Artifact, OneKitError, execute


def clean(name, data, options=None):
    return execute('text-strip-formatting', [Artifact(name, data)], options or {})


def test_plain_literal_and_encoding():
    text = '**literal** <3\r\nCaffè\u00a0ok'
    result = clean('sample.txt', text.encode('utf-16'))
    assert result.files[0].data.decode() == '**literal** <3\nCaffè ok'
    assert result.files[0].name == 'sample-plain.txt'


def test_html_omits_active_and_hidden_content():
    result = clean('sample.html', b'<head><title>omit</title></head><p>Hello <b>world</b> &amp; friends</p><script>omit()</script><div hidden>secret</div><p><a href="https://example.com">Link</a></p>')
    assert result.files[0].data.decode() == 'Hello world & friends\n\nLink'


def test_markdown_and_explicit_format():
    source = b'# Title\n\n**Bold** and [link](https://example.com)\n\n~~old~~'
    result = clean('sample.md', source)
    assert result.files[0].data.decode() == 'Title\n\nBold and link\n\nold'
    assert clean('sample.txt', b'<b>Text</b>', {'input_format': 'html'}).files[0].data == b'Text'
    assert clean('sample.txt', b'<b>Text</b>').files[0].data == b'<b>Text</b>'


def test_docx_body_and_table():
    document = Document()
    document.add_paragraph('Hello').add_run(' world').bold = True
    document.sections[0].header.paragraphs[0].text = 'Omitted header'
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'A'
    table.cell(0, 1).text = 'B'
    stream = io.BytesIO()
    document.save(stream)
    result = clean('sample.docx', stream.getvalue())
    assert result.files[0].data == b'Hello world\nA\tB'
    assert result.warnings


@pytest.mark.parametrize('data', [b'\x89PNG\r\n\x1a\n', b'a\x00b', b'\xffinvalid'])
def test_reject_binary_or_unsupported_encoding(data):
    with pytest.raises(OneKitError):
        clean('sample.txt', data)
