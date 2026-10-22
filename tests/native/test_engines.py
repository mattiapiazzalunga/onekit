# Copyright 2026 Mattia Piazzalunga
# SPDX-License-Identifier: LicenseRef-PolyForm-Shield-1.0.0
import io
import json

import pytest

from onekit import Artifact, execute
from onekit.catalog import TOOLS
from onekit.cli import main
from onekit.tools.developer import engine_for
from test_core import pdf

pytestmark = pytest.mark.native


@pytest.mark.parametrize("id", ["md-to-latex", "md-to-pdf", "latex-to-html", "latex-to-pdf", "latex-to-md", "html-to-pdf", "html-to-latex"])
def test_document_engines(id):
    source = id.split("-to-")[0]
    file = {"md": Artifact("document.md", b"# OneKit\n\nWorking document."), "latex": Artifact("document.tex", b"\\documentclass{article}\n\\begin{document}\nOneKit document.\n\\end{document}"), "html": Artifact("document.html", b"<html><body><h1>OneKit</h1><p>Working document.</p></body></html>")}[source]
    result = execute(id, [file])
    assert result.files[0].data
    if result.files[0].kind == "pdf":
        assert result.files[0].data.startswith(b"%PDF-")


@pytest.mark.parametrize("kind", ["docx", "pptx", "xlsx"])
def test_office_pdf(kind):
    stream = io.BytesIO()
    if kind == "docx":
        from docx import Document
        document = Document()
        document.add_paragraph("OneKit office conversion")
    elif kind == "pptx":
        from pptx import Presentation
        document = Presentation()
        slide = document.slides.add_slide(document.slide_layouts[0])
        slide.shapes.title.text = "OneKit office conversion"
    else:
        from openpyxl import Workbook
        document = Workbook()
        document.active.append(["OneKit", 42])
    document.save(stream)
    result = execute(f"{kind}-to-pdf", [Artifact(f"document.{kind}", stream.getvalue())])
    assert result.files[0].data.startswith(b"%PDF-")
    if kind != "xlsx":
        images = execute(f"{kind}-to-png", [Artifact(f"document.{kind}", stream.getvalue())])
        assert images.files[0].data.startswith(b"\x89PNG")


def test_audio_and_video_engines(tmp_path):
    import wave
    from onekit.tools.native import command
    stream = io.BytesIO()
    with wave.open(stream, "wb") as audio:
        audio.setnchannels(1)
        audio.setsampwidth(2)
        audio.setframerate(8000)
        audio.writeframes(b"\x00\x00" * 8000)
    converted = execute("audio-convert", [Artifact("sound.wav", stream.getvalue())], {"format": "flac"})
    assert converted.files[0].data.startswith(b"fLaC")
    command("ffmpeg", ["-nostdin", "-f", "lavfi", "-i", "color=c=blue:s=64x64:d=0.2", "-c:v", "libvpx-vp9", "input.webm"], tmp_path)
    result = execute("video-convert", [Artifact("video.webm", (tmp_path / "input.webm").read_bytes())], {"format": "mp4"})
    assert result.files[0].data[4:8] == b"ftyp"


def test_ocr_all_exports():
    from PIL import Image, ImageDraw, ImageFont
    image = Image.new("RGB", (900, 160), "white")
    font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 54)
    ImageDraw.Draw(image).text((30, 30), "OneKit document", fill="black", font=font)
    stream = io.BytesIO()
    image.save(stream, format="PNG")
    for target in ("pdf", "docx", "md", "latex", "html", "json"):
        result = execute(f"ocr-to-{target}", [Artifact("scan.png", stream.getvalue())])
        assert len(result.files[0].data) > 10


def test_metadata_engine():
    result = execute("file-metadata-set", [pdf()], {"metadata": {"Title": "OneKit"}})
    from pypdf import PdfReader
    assert PdfReader(io.BytesIO(result.files[0].data)).metadata.title == "OneKit"
    stripped = execute("file-metadata-remove", result.files)
    metadata = PdfReader(io.BytesIO(stripped.files[0].data)).metadata
    assert not (metadata or {}).get("/Title")


SAMPLES = {"javascript": "const value={a:1};", "typescript": "const value:number=1;", "jsx": "const view=<div>OneKit</div>;", "tsx": "const view=<div>OneKit</div>;", "html": "<p>OneKit</p>", "css": "body{color:blue}", "scss": "$color: blue; body { color: $color; }", "less": "@color: blue; body { color: @color; }", "markdown": "# OneKit\n\nText", "mdx": "# OneKit\n\n<Panel />", "graphql": "query { toolkit { name } }", "vue": "<template><p>OneKit</p></template>", "svelte": "<p>OneKit</p>", "astro": "---\nconst name='OneKit';\n---\n<p>{name}</p>", "solidity": "pragma solidity ^0.8.0; contract OneKit {}", "php": "<?php $name = 'OneKit';"}


@pytest.mark.parametrize("language", SAMPLES)
def test_node_format_and_validation(language):
    file = Artifact("input.txt", SAMPLES[language].encode())
    formatted = execute(f"dev-format-{language}", [file])
    assert formatted.files[0].data
    assert execute(f"dev-validate-{language}", [file]).details["valid"]


@pytest.mark.parametrize("language", ["javascript", "typescript", "jsx", "tsx", "html", "css", "svg", "graphql"])
def test_node_minifiers(language):
    text = '<svg xmlns="http://www.w3.org/2000/svg"><path d="M0 0L10 10"/></svg>' if language == "svg" else SAMPLES[language]
    assert execute(f"dev-minify-{language}", [Artifact("input.txt", text.encode())]).files[0].data


NODE_SAMPLES = {**SAMPLES, "json5": '{unquoted: 1}', "jsonc": '{"value": 1 /* comment */}', "yaml": "value: 1\n", "svg": '<svg xmlns="http://www.w3.org/2000/svg"><path d="M0 0L10 10"/></svg>'}


@pytest.mark.parametrize("tool", [tool for tool in TOOLS.values() if tool.group == "developer" and tool.id not in ['dev-format-xml', 'dev-format-svg', 'dev-format-json', 'dev-format-toml', 'dev-validate-python', 'dev-validate-xml', 'dev-validate-svg', 'dev-validate-json', 'dev-validate-json5', 'dev-validate-jsonc', 'dev-validate-yaml', 'dev-validate-toml', 'dev-validate-xsd', 'dev-validate-json-schema', 'dev-validate-csv', 'dev-minify-json'] and engine_for(tool) == "node"], ids=lambda tool: tool.id)
def test_node_developer_tools_through_cli(tool, tmp_path):
    language = tool.id.rsplit("-", 1)[-1]
    source = tmp_path / "input.txt"
    source.write_text(NODE_SAMPLES[language], encoding="utf-8")
    destination = tmp_path / "output.txt"
    assert main(["run", tool.id, str(source), "-o", str(destination)]) == 0
    assert destination.read_bytes()
    if tool.id.startswith("dev-validate-"):
        assert json.loads(destination.read_text())["valid"]


@pytest.mark.parametrize("language,filename,valid,invalid", [
    ("java", "OneKit.java", "public class OneKit {}", "public class OneKit {"),
    ("powershell", "input.ps1", "$value = 1", "function Broken {"),
    ("bash", "input.sh", "echo OneKit\n", "if then\n"),
    ("shell", "input.sh", "echo OneKit\n", "if then\n"),
])
def test_native_parsers_accept_valid_and_reject_invalid_syntax(language, filename, valid, invalid):
    for source, expected in ((valid, True), (invalid, False)):
        result = execute(f"dev-validate-{language}", [Artifact(filename, source.encode())])
        assert result.details["valid"] is expected


def test_python_formatter_cli(tmp_path):
    source = tmp_path / "input.py"
    source.write_text("value=1\n")
    output = tmp_path / "formatted.py"
    assert main(["run", "dev-format-python", str(source), "-o", str(output)]) == 0
    assert output.read_text() == "value = 1\n"
