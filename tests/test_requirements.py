# Copyright 2026 Mattia Piazzalunga
# SPDX-License-Identifier: LicenseRef-PolyForm-Shield-1.0.0
"""CLI, conversion and malformed-input regressions from the requirements audit."""
import io
import json
from pathlib import Path

import pytest
from test_core import pdf, processing_case

from onekit import Artifact, OneKitError, execute, pipeline
from onekit.catalog import TOOLS
from onekit.cli import main, read_files


def test_original_prompt_tool_requirements_remain_available():
    conversions = {
        "md": "latex json html pdf docx", "pdf": "pdfa png jpeg webp tiff docx html pptx md",
        "image": "base64 pdf pptx", "base64": "image", "docx": "pdf png md", "pptx": "pdf png",
        "latex": "html pdf md", "html": "pdf latex json md", "xlsx": "pdf csv json", "json": "xlsx", "csv": "xlsx json",
    }
    expected = {f"{source}-to-{target}" for source, targets in conversions.items() for target in targets.split()}
    expected.update("pdf-" + name for name in "merge split remove-pages organize extract-pages compress repair edit rotate sign protect unlock validate-pdfa".split())
    expected.update("image-" + name for name in "compress crop convert extract-pdf".split())
    expected.update("diff-" + name for name in "text json notebook latex pdf word code".split())
    expected.update("ocr-to-" + name for name in "pdf docx md latex html json".split())
    expected.update("file-" + name for name in "inspect metadata-set metadata-remove count encoding-convert".split())
    expected.update("security-" + name for name in ("encrypt", "decrypt"))
    expected.update(("video-convert", "audio-convert", "archive-convert", "font-convert"))
    languages = "javascript typescript jsx tsx python java c cpp csharp go rust php ruby kotlin swift dart scala lua bash shell powershell sql html css scss less xml svg json json5 jsonc yaml toml markdown mdx latex graphql vue svelte astro solidity hcl terraform dockerfile r julia elixir haskell".split()
    for operation, names in (("format", languages), ("validate", languages + "xsd json-schema csv openapi asyncapi".split()), ("minify", "javascript typescript jsx tsx html css json svg graphql".split())):
        expected.update(f"dev-{operation}-{name}" for name in names)
    assert expected <= TOOLS.keys()


@pytest.mark.parametrize("tool", [tool for tool in TOOLS.values() if tool.id in ['md-to-json', 'md-to-html', 'md-to-docx', 'pdf-to-docx', 'pdf-to-html', 'pdf-to-md', 'image-to-base64', 'image-to-pdf', 'image-to-pptx', 'base64-to-image', 'docx-to-md', 'html-to-json', 'html-to-md', 'xlsx-to-csv', 'xlsx-to-json', 'json-to-xlsx', 'csv-to-xlsx', 'csv-to-json', 'pdf-merge', 'pdf-split', 'pdf-remove-pages', 'pdf-organize', 'pdf-extract-pages', 'pdf-compress', 'pdf-repair', 'pdf-edit', 'pdf-rotate', 'pdf-protect', 'pdf-unlock', 'image-compress', 'image-crop', 'image-convert', 'image-extract-pdf', 'image-to-favicon', 'archive-convert', 'font-convert', 'file-inspect', 'file-count', 'file-encoding-convert', 'security-encrypt', 'security-decrypt', 'dev-format-xml', 'dev-format-svg', 'dev-format-json', 'dev-format-toml', 'dev-validate-python', 'dev-validate-xml', 'dev-validate-svg', 'dev-validate-json', 'dev-validate-json5', 'dev-validate-jsonc', 'dev-validate-yaml', 'dev-validate-toml', 'dev-validate-xsd', 'dev-validate-json-schema', 'dev-validate-csv', 'dev-minify-json', 'diff-text', 'diff-json', 'diff-notebook', 'diff-latex', 'diff-pdf', 'diff-word', 'diff-code', 'text-to-qrcode', 'url-to-qrcode', 'wifi-to-qrcode', 'vcard-to-qrcode', 'text-to-barcode', 'image-redact', 'pdf-watermark', 'image-effects', 'vcard-generate', 'image-upscale', 'license-pdf', 'license-word', 'license-powerpoint', 'license-excel', 'license-image', 'license-font', 'license-epub', 'license-code', 'license-dataset', 'image-rotate-flip', 'text-strip-formatting']], ids=lambda tool: tool.id)
def test_processing_cases_run_through_cli(tool, tmp_path):
    files, options = processing_case(tool)
    paths = []
    for index, file in enumerate(files):
        path = tmp_path / str(index) / file.name
        path.parent.mkdir(parents=True)
        path.write_bytes(file.data)
        paths.append(str(path))
    destination = tmp_path / "outputs"
    assert main(["run", tool.id, *paths, "--options", json.dumps(options), "-o", str(destination)]) == 0
    assert any(file.is_file() and file.stat().st_size for file in destination.rglob("*"))


@pytest.mark.parametrize("steps", [{}, [], [None], [{"tool": []}], [{"tool": "file-count", "options": []}]])
def test_malformed_pipeline_is_a_cli_error_even_with_password_prompt(steps, tmp_path, monkeypatch, capsys):
    workflow = tmp_path / "workflow.json"
    workflow.write_text(json.dumps(steps), encoding="utf-8")
    source = tmp_path / "input.txt"
    source.write_text("OneKit")
    monkeypatch.setattr("getpass.getpass", lambda *_: pytest.fail("Invalid workflows must not prompt for a password"))
    assert main(["pipeline", "--steps", str(workflow), str(source), "--ask-password", "-o", str(tmp_path / "out")]) == 2
    assert "OneKit:" in capsys.readouterr().err


def test_pipeline_preflights_all_steps_before_processing(monkeypatch):
    monkeypatch.setattr("onekit.api.execute", lambda *_args, **_kwargs: pytest.fail("Invalid workflows must be rejected first"))
    with pytest.raises(OneKitError, match="Unknown tool"):
        pipeline([{"tool": "file-count"}, {"tool": "missing"}], [Artifact("a.txt", b"a")])


def test_cli_bounds_total_input_before_reading_more_files(tmp_path, monkeypatch):
    monkeypatch.setattr("onekit.cli.MAX_BYTES", 10)
    paths = [tmp_path / name for name in ("a.txt", "b.txt")]
    for path in paths:
        path.write_bytes(b"123456")
    with pytest.raises(OneKitError, match="128 MiB"):
        read_files(paths)


@pytest.mark.parametrize("tool,text", [
    ("dev-validate-yaml", b"a: ["),
    ("dev-validate-json-schema", b'{"type":"unknown"}'),
    ("dev-validate-jsonc", b"{unquoted:1}"),
    ("dev-validate-jsonc", b'{"value":Infinity}'),
    ("dev-validate-jsonc", b'{"value":1,"value":2}'),
])
def test_invalid_structured_languages_produce_validation_reports(tool, text):
    result = execute(tool, [Artifact("input.txt", text)])
    assert result.details["valid"] is False
    assert result.details["errors"]


def test_jsonc_preserves_comment_markers_inside_strings():
    source = b'{/* comment */"url":"https://onekit.org/", "text":", } /* not a comment */", // line\n"array":[1,2,],}'
    assert execute("dev-validate-jsonc", [Artifact("input.jsonc", source)]).details["valid"]


def test_empty_tables_preserve_headers_in_excel_and_csv():
    from openpyxl import load_workbook
    source = Artifact("table.csv", b"name,value\n")
    excel = execute("csv-to-xlsx", [source]).files
    book = load_workbook(io.BytesIO(excel[0].data))
    assert list(book.active.values) == [("name", "value")]
    book.close()
    assert execute("xlsx-to-csv", excel).files[0].text().strip() == "name,value"


@pytest.mark.parametrize("text", [b'[{"value":NaN}]', b'[{"value":1,"value":2}]'])
def test_table_conversion_rejects_ambiguous_json(text):
    with pytest.raises(OneKitError):
        execute("json-to-xlsx", [Artifact("table.json", text)])


@pytest.mark.parametrize("format", ["webp", "ico", "png", "tiff"])
def test_image_slides_support_every_advertised_image_type_and_preserve_aspect(format):
    from PIL import Image
    from pptx import Presentation
    stream = io.BytesIO()
    Image.new("RGB", (32, 24), "blue").save(stream, format=format)
    result = execute("image-to-pptx", [Artifact(f"portrait.{format}", stream.getvalue())])
    presentation = Presentation(io.BytesIO(result.files[0].data))
    picture = presentation.slides[0].shapes[0]
    assert abs(picture.width / picture.height - 32 / 24) < 0.01
    assert picture.top >= 0 and picture.left >= 0
    assert picture.top + picture.height <= presentation.slide_height


def test_office_render_preserves_input_name(monkeypatch):
    def office(engine, arguments, directory, **_kwargs):
        Path(directory, "input.pdf").write_bytes(pdf().data)
    monkeypatch.setattr("onekit.tools.native.command", office)
    from docx import Document
    stream = io.BytesIO()
    document = Document()
    document.add_paragraph("Report")
    document.save(stream)
    result = execute("docx-to-png", [Artifact("report.docx", stream.getvalue())])
    assert result.files[0].name == "report.page-1.png"


def test_java_validation_preserves_public_class_filename(monkeypatch):
    calls = []
    def compiler(engine, arguments, directory, **_kwargs):
        calls.append((engine, arguments))
        assert Path(directory, "OneKit.java").read_text() == "public class OneKit {}"
        return b""
    monkeypatch.setattr("onekit.tools.native.command", compiler)
    result = execute("dev-validate-java", [Artifact("OneKit.java", b"public class OneKit {}")])
    assert result.details["valid"]
    assert calls == [("javac", ["-proc:none", "OneKit.java"])]


def test_tif_alias_is_supported():
    from PIL import Image
    stream = io.BytesIO()
    Image.new("RGB", (10, 10)).save(stream, format="TIFF")
    assert execute("image-convert", [Artifact("image.tif", stream.getvalue())], {"format": "png"}).files[0].data.startswith(b"\x89PNG")


@pytest.mark.parametrize("target", ["png", "jpeg", "webp", "tiff", "pptx"])
def test_pdf_renders_and_slides_have_real_outputs(target):
    from PIL import Image
    result = execute(f"pdf-to-{target}", [pdf()], {"dpi": 72})
    if target == "pptx":
        from pptx import Presentation
        assert len(Presentation(io.BytesIO(result.files[0].data)).slides) == 2
    else:
        assert len(result.files) == 2
        for file in result.files:
            with Image.open(io.BytesIO(file.data)) as image:
                assert image.width > 0 and image.height > 0


@pytest.mark.parametrize("data,mime", [
    (b"RIFF\x00\x00\x00\x00WAVE", "audio/wav"),
    (b"RIFF\x00\x00\x00\x00AVI ", "video/x-msvideo"),
    (b"\x00\x00\x00\x18ftypisom", "video/mp4"),
    (b"\x00\x01\x00\x00binary", "font/ttf"),
    (b"OTTO\x00\x00\x00\x00", "font/otf"),
    (b"ID3\x04\x00\x00", "audio/mpeg"),
    (b"\x00\x01\x00unknown", "application/octet-stream"),
])
def test_inspector_identifies_binary_signatures_without_guessing_text(data, mime):
    details = execute("file-inspect", [Artifact("unknown.bin", data)]).details
    assert details["mime"] == mime
    assert details["encoding"] is None


def test_tool_directories_match_runtime_catalog():
    from dataclasses import asdict
    root = Path(__file__).resolve().parents[1]
    for tool in TOOLS.values():
        metadata = json.loads((root / "tools" / tool.group / tool.id / "tool.json").read_text(encoding="utf-8"))
        assert {key: metadata[key] for key in asdict(tool)} == json.loads(json.dumps(asdict(tool)))


def test_decryption_can_restore_a_downloadable_zip_with_folder_paths():
    import zipfile
    files = [Artifact("private/a.txt", b"OneKit"), Artifact("private/nested/b.txt", b"folder data")]
    options = {"password": "a-long-private-password"}
    encrypted = execute("security-encrypt", files, options)
    restored = execute("security-decrypt", encrypted.files, {**options, "restore": "zip"})
    with zipfile.ZipFile(io.BytesIO(restored.files[0].data)) as archive:
        assert archive.namelist() == [file.name for file in files]
        assert [archive.read(file.name) for file in files] == [file.data for file in files]


def test_windows_shell_uses_git_instead_of_the_wsl_launcher(tmp_path, monkeypatch):
    from onekit.tools.native import windows_shell
    monkeypatch.setenv("SystemRoot", str(tmp_path / "Windows"))
    monkeypatch.setenv("ProgramFiles", str(tmp_path / "Programs"))
    monkeypatch.delenv("LOCALAPPDATA", raising=False)
    launcher = str(tmp_path / "Windows" / "System32" / "bash.exe")
    with pytest.raises(OneKitError, match="WSL launcher"):
        windows_shell("bash", launcher)
    bash = tmp_path / "Programs" / "Git" / "usr" / "bin" / "bash.exe"
    bash.parent.mkdir(parents=True)
    bash.touch()
    assert windows_shell("bash", launcher) == str(bash)
    custom = str(tmp_path / "custom" / "bash.exe")
    assert windows_shell("bash", custom) == custom


def test_doctor_requires_the_powershell_formatter_module(monkeypatch):
    from onekit.cli import doctor
    monkeypatch.setattr("onekit.tools.native.executable", lambda _name: "pwsh")
    def unavailable(*_args, **_kwargs):
        raise OneKitError("Module unavailable")
    monkeypatch.setattr("onekit.tools.native.command", unavailable)
    report = doctor("dev-format-powershell")[0]
    assert not report["ready"]
    assert "PSScriptAnalyzer" in report["missing"][0]
