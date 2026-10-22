# Copyright 2026 Mattia Piazzalunga
# SPDX-License-Identifier: LicenseRef-PolyForm-Shield-1.0.0
import base64
import io
import json
import zipfile

import pytest

from onekit import Artifact, OneKitError, execute, pipeline
from onekit.catalog import TOOLS


def png():
    from PIL import Image
    stream = io.BytesIO()
    Image.new("RGB", (32, 24), (51, 102, 204)).save(stream, format="PNG")
    return Artifact("image.png", stream.getvalue())


def pdf():
    from reportlab.pdfgen import canvas
    stream = io.BytesIO()
    document = canvas.Canvas(stream)
    document.drawString(30, 700, "OneKit regression document")
    document.drawImage(__import__('reportlab').lib.utils.ImageReader(io.BytesIO(png().data)), 30, 600, 32, 24)
    document.showPage()
    document.drawString(30, 700, "Second page")
    document.save()
    return Artifact("document.pdf", stream.getvalue())


def processing_case(tool):
    id = tool.id
    options = {}
    if id.startswith("license-"):
        from pathlib import Path

        from test_license_check import samples
        kind = id.removeprefix("license-")
        if kind in ("audio", "video"):
            import wave
            if kind == "audio":
                stream = io.BytesIO()
                with wave.open(stream, "wb") as audio:
                    audio.setnchannels(1)
                    audio.setsampwidth(2)
                    audio.setframerate(8000)
                    audio.writeframes(b"\0" * 1600)
                name, data = "sample.wav", stream.getvalue()
            else:
                name, data = "sample.mp4", (Path(__file__).parent / "fixtures" / "license-sample.mp4").read_bytes()
            return [Artifact(name, data)], {"_media_metadata": {"streams": [{"codec_type": kind}], "format": {}}}
        name, data = samples()[kind]
        return [Artifact(name, data)], options
    if id == "text-strip-formatting":
        return [Artifact("sample.html", b"<p>Hello <strong>OneKit</strong></p><script>never_run()</script>")], {}
    if id == "image-rotate-flip":
        return [png()], {"angle": "90", "flip": "horizontal"}
    if id == "image-effects":
        return [png()], {"effect": "pixelate", "scope": "all", "strength": 4}
    if id == "pdf-watermark":
        return [pdf()], {"text": "OneKit draft", "pages": "1", "opacity": .25}
    if id == "image-redact":
        return [png()], {"regions": [{"x": .1, "y": .1, "width": .5, "height": .5}]}
    if id == "text-to-barcode":
        return [Artifact("input.txt", b"OneKit-123")], options
    if id == "text-to-qrcode":
        return [Artifact("input.txt", "OneKit QR — ciao!".encode())], options
    if id == "url-to-qrcode":
        return [Artifact("input.txt", b"https://example.com/?hello=world")], options
    if id == "wifi-to-qrcode":
        return [Artifact("input.json", b'{"ssid":"OneKit","password":"test-password","security":"WPA","hidden":true}')], options
    if id in ("vcard-to-qrcode", "vcard-generate"):
        return [Artifact("input.json", b'{"first_name":"Ada","last_name":"Lovelace","phones":[{"value":"+39012345678","type":"cell"}]}')], options
    if id in ("diff-pdf",):
        return [pdf(), pdf()], options
    if id in ("diff-word", "docx-to-md"):
        from docx import Document
        document = Document()
        document.add_paragraph("OneKit works")
        stream = io.BytesIO()
        document.save(stream)
        return [Artifact("document.docx", stream.getvalue())] * (2 if id == "diff-word" else 1), options
    if id == "archive-convert":
        from onekit.tools.archives import pack
        return [Artifact("files.zip", pack([Artifact("a.txt", b"OneKit")]))], {"format": "tar.gz"}
    if id == "font-convert":
        from fontTools.fontBuilder import FontBuilder
        from fontTools.pens.ttGlyphPen import TTGlyphPen
        builder = FontBuilder(1000, isTTF=True)
        builder.setupGlyphOrder([".notdef", "space"])
        builder.setupCharacterMap({32: "space"})
        builder.setupGlyf({name: TTGlyphPen(None).glyph() for name in (".notdef", "space")})
        builder.setupHorizontalMetrics({".notdef": (500, 0), "space": (500, 0)})
        builder.setupHorizontalHeader(ascent=800, descent=-200)
        builder.setupNameTable({"familyName": "OneKitTest", "styleName": "Regular"})
        builder.setupOS2(sTypoAscender=800, sTypoDescender=-200, usWinAscent=800, usWinDescent=200)
        builder.setupPost()
        stream = io.BytesIO()
        builder.save(stream)
        return [Artifact("font.ttf", stream.getvalue())], {"format": "woff"}
    if id == "security-encrypt":
        return [Artifact("private.txt", b"private test data")], {"password": "a-long-test-password"}
    if id == "security-decrypt":
        options = {"password": "a-long-test-password"}
        return execute("security-encrypt", [Artifact("private.txt", b"private test data")], options).files, options
    if id.startswith("diff-"):
        if id == "diff-json":
            return [Artifact("a.json", b'{"a":1}'), Artifact("b.json", b'{"a":2}')], options
        if id == "diff-notebook":
            return [Artifact("a.ipynb", json.dumps({"cells": [{"cell_type": "code", "source": [text]}]}).encode()) for text in ("print(1)", "print(2)")], options
        return [Artifact("a.txt", b"one\ntwo\n"), Artifact("b.txt", b"one\nthree\n")], options
    if id.startswith("pdf-") or id == "image-extract-pdf":
        if id in ("pdf-to-md", "pdf-to-html", "pdf-to-docx"):
            return [pdf()], options
        files = [pdf()]
        if id == "pdf-merge":
            files.append(Artifact("second.pdf", pdf().data))
        if id in ("pdf-remove-pages", "pdf-organize", "pdf-extract-pages"):
            options["pages"] = "1"
        if id == "pdf-edit":
            options["text"] = "Added by OneKit"
        if id == "pdf-protect":
            options["password"] = "test-password"
        if id == "pdf-unlock":
            options["password"] = "test-password"
            files = execute("pdf-protect", files, options).files
        return files, options
    if id.startswith("image-"):
        if id == "image-crop":
            options["box"] = [1, 1, 10, 10]
        return [png()], options
    if id == "base64-to-image":
        return [Artifact("image.base64", base64.b64encode(png().data))], options
    if id.startswith("xlsx-"):
        return execute("json-to-xlsx", [Artifact("rows.json", b'[{"name":"OneKit","value":1}]')]).files, options
    if id.startswith("csv-"):
        return [Artifact("rows.csv", b"name,value\nOneKit,1\n")], options
    if id == "json-to-xlsx":
        return [Artifact("rows.json", b'[{"name":"OneKit","value":1}]')], options
    if id.startswith("md-"):
        return [Artifact("text.md", b"# OneKit\n\n**Works** locally.")], options
    if id.startswith("html-"):
        return [Artifact("text.html", b"<h1>OneKit</h1><p>Works locally.</p>")], options
    if id.startswith("dev-"):
        language = id.rsplit("-", 1)[1]
        if id == "dev-validate-json-schema":
            language = "json-schema"
        value = {"json": '{"b":2,"a":1}', "json5": '{"b":2,"a":1}', "jsonc": '{"a":1 /* comment */}', "json-schema": '{"type":"object"}', "yaml": 'name: OneKit', "xsd": '<xs:schema xmlns:xs="http://www.w3.org/2001/XMLSchema"><xs:element name="item" type="xs:string"/></xs:schema>', "python": "x = 1\n", "xml": "<root/>", "svg": '<svg xmlns="http://www.w3.org/2000/svg"/>', "toml": 'name = "OneKit"', "csv": "a,b\n1,2"}[language]
        extension = {"python": "py", "json-schema": "json"}.get(language, language)
        return [Artifact(f"input.{extension}", value.encode())], options
    return [Artifact("text.txt", b"OneKit works\non your device.\n")], options


@pytest.mark.parametrize("tool", [t for t in TOOLS.values() if t.id in ['md-to-json', 'md-to-html', 'md-to-docx', 'pdf-to-docx', 'pdf-to-html', 'pdf-to-md', 'image-to-base64', 'image-to-pdf', 'image-to-pptx', 'base64-to-image', 'docx-to-md', 'html-to-json', 'html-to-md', 'xlsx-to-csv', 'xlsx-to-json', 'json-to-xlsx', 'csv-to-xlsx', 'csv-to-json', 'pdf-merge', 'pdf-split', 'pdf-remove-pages', 'pdf-organize', 'pdf-extract-pages', 'pdf-compress', 'pdf-repair', 'pdf-edit', 'pdf-rotate', 'pdf-protect', 'pdf-unlock', 'image-compress', 'image-crop', 'image-convert', 'image-extract-pdf', 'image-to-favicon', 'archive-convert', 'font-convert', 'file-inspect', 'file-count', 'file-encoding-convert', 'security-encrypt', 'security-decrypt', 'dev-format-xml', 'dev-format-svg', 'dev-format-json', 'dev-format-toml', 'dev-validate-python', 'dev-validate-xml', 'dev-validate-svg', 'dev-validate-json', 'dev-validate-json5', 'dev-validate-jsonc', 'dev-validate-yaml', 'dev-validate-toml', 'dev-validate-xsd', 'dev-validate-json-schema', 'dev-validate-csv', 'dev-minify-json', 'diff-text', 'diff-json', 'diff-notebook', 'diff-latex', 'diff-pdf', 'diff-word', 'diff-code', 'text-to-qrcode', 'url-to-qrcode', 'wifi-to-qrcode', 'vcard-to-qrcode', 'text-to-barcode', 'image-redact', 'pdf-watermark', 'image-effects', 'vcard-generate', 'image-upscale', 'license-pdf', 'license-word', 'license-powerpoint', 'license-excel', 'license-image', 'license-font', 'license-epub', 'license-code', 'license-dataset', 'image-rotate-flip', 'text-strip-formatting']], ids=lambda t: t.id)
def test_processing_contract(tool):
    files, options = processing_case(tool)
    result = execute(tool.id, files, options)
    assert result.files
    assert all(file.data for file in result.files)


def test_semantic_json_diff():
    result = execute("diff-json", [Artifact("a.json", b'{"a":1,"b":2}'), Artifact("b.json", b'{"b":2,"a":1}')])
    assert result.details["equal"] is True


def test_notebook_filter():
    def notebook(code):
        return Artifact("n.ipynb", json.dumps({"cells": [{"cell_type": "markdown", "source": ["same"]}, {"cell_type": "code", "source": [code]}]}).encode())
    assert execute("diff-notebook", [notebook("x=1"), notebook("x=2")], {"mode": "markdown"}).details["equal"]


def test_pipeline_types_and_trace():
    result = pipeline([{"tool": "md-to-html"}, {"tool": "html-to-md"}], [Artifact("a.md", b"# OneKit")])
    assert "OneKit" in result.files[0].text()
    assert len(result.details["steps"]) == 2
    with pytest.raises(OneKitError, match="Expected file types"):
        pipeline([{"tool": "md-to-html"}, {"tool": "pdf-split"}], [Artifact("a.md", b"# OneKit")])


def test_xlsx_formula_injection():
    from openpyxl import load_workbook
    result = execute("json-to-xlsx", [Artifact("rows.json", b'[{"value":"=HYPERLINK(1)"}]')])
    book = load_workbook(io.BytesIO(result.files[0].data))
    assert book.active["A2"].data_type == "s"


@pytest.mark.parametrize("name", ["../secret", "/secret", "C:\\secret", "folder/../../secret"])
def test_unsafe_names(name):
    with pytest.raises(OneKitError):
        Artifact(name, b"content")


def test_archive_traversal_and_roundtrip():
    from onekit.tools.archives import pack, unpack
    files = [Artifact("folder/data.txt", b"safe")]
    for format in ("zip", "tar", "tar.gz", "tar.bz2", "tar.xz"):
        assert unpack(pack(files, format)) == files
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w") as archive:
        archive.writestr("../escape", "no")
    with pytest.raises(OneKitError, match="unsafe path"):
        unpack(stream.getvalue())


def test_encryption_authentication():
    files = [Artifact("folder/private.txt", b"private content")]
    options = {"password": "a-long-test-password"}
    encrypted = execute("security-encrypt", files, options)
    assert files[0].data not in encrypted.files[0].data
    assert execute("security-decrypt", encrypted.files, options).files == files
    with pytest.raises(OneKitError, match="Wrong password"):
        execute("security-decrypt", encrypted.files, {"password": "a-different-password"})
    data = bytearray(encrypted.files[0].data)
    import struct
    name_size, extra_size = struct.unpack_from("<HH", data, 26)
    data[30 + name_size + extra_size + 18] ^= 1
    corrupted = Artifact("bad.zip", bytes(data))
    with pytest.raises(OneKitError, match="damaged"):
        execute("security-decrypt", [corrupted], options)


def test_invalid_input_reports():
    result = execute("dev-validate-json", [Artifact("a.json", b"{invalid}")])
    assert result.details["valid"] is False
    with pytest.raises(OneKitError, match="UTF-8"):
        execute("file-count", [Artifact("a.bin", b"\xff\xfe\x00")])
    with pytest.raises(OneKitError, match="one input"):
        execute("file-count", [Artifact("a.txt", b"a"), Artifact("b.txt", b"b")])


def test_pdf_pages_and_rotation():
    from pypdf import PdfReader
    result = execute("pdf-organize", [pdf()], {"pages": "2,1,2"})
    reader = PdfReader(io.BytesIO(result.files[0].data))
    assert len(reader.pages) == 3
    assert "Second page" in reader.pages[0].extract_text()
    rotated = execute("pdf-rotate", result.files, {"angle": 90})
    assert PdfReader(io.BytesIO(rotated.files[0].data)).pages[0].rotation == 90


def test_cli_preserves_existing_outputs(tmp_path):
    from onekit.cli import main
    source = tmp_path / "source.txt"
    source.write_text("OneKit", encoding="utf-8")
    dest = tmp_path / "count.json"
    assert main(["run", "file-count", str(source), "-o", str(dest)]) == 0
    original = dest.read_bytes()
    assert main(["run", "file-count", str(source), "-o", str(dest)]) == 2
    assert dest.read_bytes() == original


def test_strict_users_can_disable_external_copyleft(monkeypatch):
    from onekit.tools.native import executable
    monkeypatch.setenv("ONEKIT_ALLOW_COPYLEFT", "0")
    with pytest.raises(OneKitError, match="disabled"):
        executable("ffmpeg")




def test_python_validation_never_executes(tmp_path):
    marker = tmp_path / "executed"
    code = f"from pathlib import Path\nPath({str(marker)!r}).touch()\n"
    assert execute("dev-validate-python", [Artifact("a.py", code.encode())]).details["valid"]
    assert not marker.exists()


@pytest.mark.parametrize("text", [b"NaN", b"Infinity", b'{"a":1,"a":2}'])
def test_standard_json_validation(text):
    assert not execute("dev-validate-json", [Artifact("input.json", text)]).details["valid"]


def test_pipeline_maps_multiple_outputs():
    result = pipeline([{"tool": "pdf-split"}, {"tool": "pdf-rotate", "options": {"angle": 90}}], [pdf()])
    assert len(result.files) == 2


def test_markdown_word_structure():
    from docx import Document
    result = execute("md-to-docx", [Artifact("document.md", b"# Heading\n\nParagraph with **bold**.\n\n## Next\n\nLast.")])
    document = Document(io.BytesIO(result.files[0].data))
    assert document.paragraphs[0].style.name == "Heading 1"
    assert document.paragraphs[2].style.name == "Heading 2"
    assert any(run.bold for paragraph in document.paragraphs for run in paragraph.runs)


def test_notebook_merge_keeps_notebook_and_unselected_cells():
    left = {"nbformat": 4, "metadata": {"name": "left"}, "cells": [{"cell_type": "markdown", "source": ["Changed"], "metadata": {}}, {"cell_type": "code", "source": ["x=1"], "outputs": [], "metadata": {}}]}
    right = {"nbformat": 4, "metadata": {"name": "right"}, "cells": [{"cell_type": "markdown", "source": ["Original"], "metadata": {}}, {"cell_type": "code", "source": ["x=2"], "outputs": [], "metadata": {}}]}
    files = [Artifact("left.ipynb", json.dumps(left).encode()), Artifact("right.ipynb", json.dumps(right).encode())]
    result = execute("diff-notebook", files, {"mode": "markdown", "merge": {"from": "left", "block": 0}})
    merged = json.loads(result.details["sources"]["right"])
    assert merged["cells"][0]["source"] == ["Changed"]
    assert merged["cells"][1]["source"] == ["x=2"]
    assert merged["metadata"] == right["metadata"]
    assert result.details["equal"]
