# Copyright 2026 Mattia Piazzalunga
# SPDX-License-Identifier: LicenseRef-PolyForm-Shield-1.0.0
import csv
import io
import json
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook
from pypdf import PdfReader

from onekit import Artifact, OneKitError, execute, pipeline

EXPORTS = ['md','html','docx','pptx','xlsx','csv','json','xml','yaml','epub']
IMPORTS = ['pdf','docx','html','md','csv','xlsx','json','xml','yaml']
IDS = [source+'-to-txt' for source in EXPORTS]+['txt-to-'+target for target in IMPORTS]


def samples(tool):
    source,target = tool.id.split('-to-')
    if source=='txt':
        text = {'json':'{"name":"Caffè","count":2,"active":true}',
                'xml':'<?xml version="1.0" encoding="utf-8"?><rows><row>Ada</row></rows>',
                'yaml':'# retained comment\nname: Ada\ncount: 2\n',
                'csv':'name\tcount\nAda\t00123\n','xlsx':'name\tcount\nAda\t00123\n'}.get(target,'Hello OneKit\nCaffè & literal <tags>.')
        return [Artifact('sample.txt',text.encode())],{}
    if source=='md':
        return [Artifact('sample.md',b'# Heading\n\nHello **OneKit**.\n\n<script>hidden</script>')],{}
    if source=='html':
        return [Artifact('sample.html',b'<html><head><title>hidden</title></head><body><h1>Heading</h1><p>Hello <b>OneKit</b>.</p><p hidden>hidden</p><script>hidden</script></body></html>')],{}
    if source=='docx':
        from docx import Document
        document = Document();document.add_paragraph('Hello OneKit')
        document.sections[0].header.paragraphs[0].text = 'Header excluded'
        table = document.add_table(rows=1,cols=2);table.cell(0,0).text='Ada';table.cell(0,1).text='00123'
        stream = io.BytesIO();document.save(stream)
        return [Artifact('sample.docx',stream.getvalue())],{}
    if source=='pptx':
        from pptx import Presentation
        from pptx.util import Inches
        book = Presentation()
        for text in ['First slide','Second slide']:
            slide = book.slides.add_slide(book.slide_layouts[6]);slide.shapes.add_textbox(Inches(1),Inches(1),Inches(5),Inches(1)).text = text
        order = book.slides._sldIdLst;order.append(order[0])
        stream = io.BytesIO();book.save(stream)
        return [Artifact('slides.pptx',stream.getvalue())],{}
    if source=='xlsx':
        book = Workbook();book.active.append(['name','count']);book.active.append(['Ada','00123'])
        stream = io.BytesIO();book.save(stream)
        return [Artifact('sample.xlsx',stream.getvalue())],{}
    if source=='csv':
        return [Artifact('sample.csv',b'name,note\nAda,"comma, tab\there"\n')],{}
    if source=='json':
        return [Artifact('sample.json',b'{"name":"Ada","count":2,"empty":null}')],{}
    if source=='xml':
        return [Artifact('sample.xml',b'<rows><row name="Ada">2</row></rows>')],{}
    if source=='yaml':
        return [Artifact('sample.yaml',b'# retained comment\nname: Ada\ncount: 2\n')],{}
    from test_new_converters import samples as converter_samples
    return converter_samples(SimpleNamespace(id='epub-to-pdf'))


@pytest.mark.parametrize('identifier',IDS)
def test_each_text_converter_exports_a_real_compatible_artifact(identifier):
    from onekit.catalog import TOOLS
    files,options = samples(TOOLS[identifier]);result = execute(identifier,files,options)
    assert len(result.files)==1 and result.files[0].data
    source,target = identifier.split('-to-')
    assert result.files[0].kind==target
    if identifier=='txt-to-pdf':
        text = '\n'.join(page.extract_text() for page in PdfReader(io.BytesIO(result.files[0].data)).pages)
        assert 'Hello OneKit' in text and 'Caffè' in text
    if source in ('html','md'):
        assert 'Hello OneKit' in result.files[0].text() and 'hidden' not in result.files[0].text()
    if source=='docx':
        assert 'Ada\t00123' in result.files[0].text() and 'Header excluded' not in result.files[0].text()
    if source=='pptx':
        assert result.files[0].text().index('Second slide')<result.files[0].text().index('First slide')
    if source=='epub':
        assert result.files[0].text().index('First chapter')<result.files[0].text().index('Second chapter')
    if identifier=='txt-to-xlsx':
        assert load_workbook(io.BytesIO(result.files[0].data)).active['B2'].value=='00123'


@pytest.mark.parametrize('kind',['html','md','docx'])
def test_plain_text_roundtrip_preserves_literal_markup(kind):
    text = '# Heading\nLiteral <tags>, **stars** and [links](url).'
    result = pipeline([{'tool':'txt-to-'+kind},{'tool':kind+'-to-txt'}],[Artifact('literal.txt',text.encode())])
    assert result.files[0].text()==text


@pytest.mark.parametrize('kind',['json','xml','yaml'])
def test_structured_roundtrips_validate_content_without_inventing_schema(kind):
    from onekit.catalog import TOOLS
    files,_ = samples(TOOLS[kind+'-to-txt'])
    result = pipeline([{'tool':kind+'-to-txt'},{'tool':'txt-to-'+kind}],files)
    if kind=='json':
        assert json.loads(result.files[0].text())==json.loads(files[0].text())
    else:
        assert result.files[0].text()==files[0].text()


@pytest.mark.parametrize('kind',['json','xml'])
def test_arbitrary_prose_cannot_be_exported_as_structured_data(kind):
    with pytest.raises(OneKitError):
        execute('txt-to-'+kind,[Artifact('prose.txt',b'Hello world, this is prose.')])


def test_utf16_text_is_decoded_before_json_validation():
    result = execute('txt-to-json',[Artifact('data.txt','{"name":"Caffè"}'.encode('utf-16'))],{'source_encoding':'utf-16'})
    assert json.loads(result.files[0].text())=={'name':'Caffè'}


def test_csv_text_roundtrip_preserves_quoted_tabs_and_newlines():
    text = 'name,note\r\nAda,"tab\there, comma\nand newline"\r\n'
    result = pipeline([{'tool':'csv-to-txt'},{'tool':'txt-to-csv'}],[Artifact('table.csv',text.encode())])
    assert list(csv.reader(io.StringIO(result.files[0].text())))==list(csv.reader(io.StringIO(text)))


def test_txt_spreadsheet_preserves_leading_zeros_and_formula_strings():
    result = execute('txt-to-xlsx',[Artifact('data.txt',b'code\tnote\n00123\t=1+1\n')])
    sheet = load_workbook(io.BytesIO(result.files[0].data)).active
    assert sheet['A2'].value=='00123' and sheet['B2'].value=='=1+1' and sheet['B2'].data_type=='s'
    with pytest.raises(OneKitError,match='column counts'):
        execute('txt-to-xlsx',[Artifact('data.txt',b'a\tb\n1\t2\t3\n')])


def test_invalid_xml_entities_and_unsafe_yaml_remain_rejected():
    with pytest.raises(OneKitError):
        execute('txt-to-xml',[Artifact('bad.txt',b'<!DOCTYPE a [<!ENTITY x "expanded">]><a>&x;</a>')])
    with pytest.raises(OneKitError):
        execute('txt-to-yaml',[Artifact('bad.txt',b'!!python/object/apply:os.system [echo bad]')])


def test_markdown_literal_block_preserves_entities_indentation_and_embedded_fences():
    text = '    *literal* &amp; <script>text</script>\n```code```\n\tTabbed line'
    result = pipeline([{'tool':'txt-to-md'},{'tool':'md-to-txt'}],[Artifact('literal.txt',text.encode())])
    assert result.files[0].text()==text
