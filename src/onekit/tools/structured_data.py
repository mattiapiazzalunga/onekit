# Copyright 2026 Mattia Piazzalunga
# SPDX-License-Identifier: LicenseRef-PolyForm-Shield-1.0.0
"""Bounded XML, JSON, YAML and tabular interchange."""
import csv
import datetime
import io
import json
import math
import re
from xml.etree import ElementTree as ET

from ..models import Artifact, OneKitError, Result, output, parse_json

_INVALID_XML_TEXT = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]')


def bounded(value, depth=0, active=None, count=None):
    active = set() if active is None else active
    count = [0] if count is None else count
    count[0] += 1
    if depth > 48 or count[0] > 100000:
        raise OneKitError('Document is too complex; reduce nesting or split the data.')
    if isinstance(value, (dict, list)):
        if id(value) in active:
            raise OneKitError('Cyclic YAML aliases cannot be converted to JSON.')
        active.add(id(value))
        if isinstance(value, dict):
            if any(not isinstance(key, str) for key in value):
                raise OneKitError('JSON object keys must be strings.')
            result = {key: bounded(part, depth+1, active, count) for key, part in value.items()}
        else:
            result = [bounded(part, depth+1, active, count) for part in value]
        active.remove(id(value))
        return result
    if isinstance(value, (datetime.date, datetime.datetime)):
        return value.isoformat()
    if value is None or isinstance(value, (str, bool, int)) or isinstance(value, float) and math.isfinite(value):
        return value
    raise OneKitError('Use JSON-compatible scalar values; binary, sets and non-finite numbers are not supported.')


def xml_root(data):
    from defusedxml import ElementTree
    try:
        root = ElementTree.fromstring(data, forbid_dtd=True, forbid_entities=True, forbid_external=True)
    except Exception as error:
        raise OneKitError('Invalid XML. DTDs and entity declarations are not supported.') from error
    pending, count = [(root,0)], 0
    while pending:
        node, depth = pending.pop(); count += 1
        if depth > 48 or count > 100000:
            raise OneKitError('Document is too complex; reduce XML nesting or split the data.')
        pending.extend((child,depth+1) for child in node)
    return root


def xml_value(node):
    value = {'@'+key: text for key,text in node.attrib.items()}
    children = list(node)
    if not children and not value:
        return node.text or ''
    for child in children:
        part = xml_value(child)
        if child.tag in value:
            if not isinstance(value[child.tag],list):
                value[child.tag] = [value[child.tag]]
            value[child.tag].append(part)
        else:
            value[child.tag] = part
    if (node.text or '').strip():
        value['#text'] = node.text.strip()
    if any((child.tail or '').strip() for child in children):
        raise OneKitError('Mixed XML content is not tabular data. Remove text between child elements before conversion.')
    return value


def xml_text(value):
    if _INVALID_XML_TEXT.search(value):
        raise OneKitError('XML cannot represent control characters or invalid Unicode in keys or values.')
    return value


def json_node(value, tag='item', name=None):
    node = ET.Element(tag)
    if name is not None:
        node.set('name',xml_text(name))
    if isinstance(value,dict):
        node.set('type','object')
        node.extend([json_node(part,name=key) for key,part in value.items()])
    elif isinstance(value,list):
        node.set('type','array');node.extend([json_node(part) for part in value])
    elif value is None:
        node.set('type','null')
    elif isinstance(value,bool):
        node.set('type','boolean');node.text = 'true' if value else 'false'
    elif isinstance(value,(float,int)):
        node.set('type','number');node.text = json.dumps(value,allow_nan=False)
    else:
        node.set('type','string');node.text = xml_text(value)
    return node


def node_json(node):
    kind = node.get('type')
    if kind == 'object':
        result = {}
        for child in node:
            key = child.get('name')
            if key is None or key in result:
                raise OneKitError('Typed XML object keys must be present and unique.')
            result[key] = node_json(child)
        return result
    if kind == 'array':
        return [node_json(child) for child in node]
    if kind == 'null':
        return None
    text = node.text or ''
    if kind == 'boolean' and text in ('true','false'):
        return text == 'true'
    if kind == 'number':
        value = parse_json(text)
        if isinstance(value,(int,float)) and not isinstance(value,bool):
            return value
    if kind == 'string':
        return text
    raise OneKitError('Invalid typed XML value.')


def flat_records(root):
    children = list(root)
    if not children:
        raise OneKitError('XML tables require repeated record elements under a root element.')
    if len({child.tag for child in children}) != 1:
        raise OneKitError('XML tables require records with the same element name.')
    rows = []
    for record in children:
        row = dict(record.attrib)
        for cell in record:
            if list(cell):
                raise OneKitError('XML tables require flat fields. Convert nested XML to JSON first.')
            name = cell.get('name') if cell.tag == 'cell' else cell.tag
            if not name or name in row:
                raise OneKitError('XML record field names must be present and unique.')
            row[name] = cell.text or ''
        rows.append(row)
    return rows


def table_rows(file, source, options):
    from .converters import bounded_rows
    if source == 'xml':
        return flat_records(xml_root(file.data))
    if source == 'csv':
        delimiter = str(options.get('delimiter',','))
        if len(delimiter) != 1 or delimiter in '\r\n':
            raise OneKitError('Choose a single-character CSV delimiter.')
        reader = csv.DictReader(io.StringIO(file.text()),delimiter=delimiter)
        if not reader.fieldnames or len(set(reader.fieldnames)) != len(reader.fieldnames):
            raise OneKitError('CSV column headers must be present and unique.')
        rows = bounded_rows(reader)
        if any(None in row for row in rows):
            raise OneKitError('CSV rows contain more values than their header.')
        return rows
    if source == 'json':
        value = bounded(parse_json(file.text()))
        if not isinstance(value,list) or any(not isinstance(row,dict) for row in value):
            raise OneKitError('Expected a JSON array of objects for CSV conversion.')
        return value
    from openpyxl import load_workbook
    book = load_workbook(io.BytesIO(file.data),read_only=True,data_only=True)
    try:
        sheet = book[options['sheet']] if options.get('sheet') else book.active
        values = sheet.iter_rows(values_only=True)
        keys = [str(value) if value is not None else f'Column {index+1}' for index,value in enumerate(next(values,[]))]
        if not keys or len(set(keys)) != len(keys):
            raise OneKitError('Spreadsheet column headers must be present and unique.')
        return bounded_rows(dict(zip(keys,row)) for row in values)
    finally:
        book.close()


def scalar(value):
    if value is None:
        return ''
    if isinstance(value,(dict,list)):
        return json.dumps(value,ensure_ascii=False,allow_nan=False)
    return str(value)


def run(tool, files, options):
    file = files[0];source,target = tool.id.split('-to-')
    if source == 'yaml':
        import yaml
        class Loader(yaml.SafeLoader):
            pass
        def mapping(loader,node):
            loader.flatten_mapping(node)
            result = {}
            for key_node,value_node in node.value:
                key = loader.construct_object(key_node)
                if not isinstance(key,str) or key in result:
                    raise OneKitError('YAML mapping keys must be unique strings.')
                result[key] = loader.construct_object(value_node)
            return result
        Loader.add_constructor(yaml.resolver.BaseResolver.DEFAULT_MAPPING_TAG,mapping)
        value = bounded(yaml.load(file.text(),Loader=Loader))
        return Result([output(file,'json',json.dumps(value,ensure_ascii=False,allow_nan=False,indent=2))])
    if source == 'json' and target in ('xml','yaml'):
        value = bounded(parse_json(file.text()))
        if target == 'yaml':
            import yaml
            data = yaml.safe_dump(value,allow_unicode=True,sort_keys=False)
        else:
            root = json_node(value,'root');root.set('data-onekit','json-v1')
            data = ET.tostring(root,encoding='unicode',xml_declaration=True)
        return Result([output(file,target,data)])
    if source == 'xml' and target == 'json':
        root = xml_root(file.data)
        value = node_json(root) if root.get('data-onekit') == 'json-v1' else {root.tag:xml_value(root)}
        return Result([output(file,'json',json.dumps(bounded(value),ensure_ascii=False,indent=2))])
    rows = table_rows(file,source,options)
    if len(rows) > 100000:
        raise OneKitError('Tables may contain at most 100,000 rows.')
    keys = list(dict.fromkeys(key for row in rows for key in row))
    if len(keys) > 1000:
        raise OneKitError('Tables may contain at most 1,000 columns.')
    if target == 'xml':
        root = ET.Element('rows')
        for row in rows:
            record = ET.SubElement(root,'row')
            for key in keys:
                ET.SubElement(record,'cell',name=xml_text(key)).text = xml_text(scalar(row.get(key)))
        return Result([output(file,'xml',ET.tostring(root,encoding='unicode',xml_declaration=True))])
    # Delegate tabular encoding to the existing shared spreadsheet writer.
    from .converters import tables
    if target == 'csv':
        rows = [{key:scalar(value) if isinstance(value,(dict,list)) else value for key,value in row.items()} for row in rows]
    proxy = Artifact(file.stem+'.json',json.dumps(rows,ensure_ascii=False,default=str).encode())
    return tables(proxy,'json',target,options)
