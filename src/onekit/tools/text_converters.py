# Copyright 2026 Mattia Piazzalunga
# SPDX-License-Identifier: LicenseRef-PolyForm-Shield-1.0.0
"""Plain text extraction and compatible, explicit conversions from TXT."""
import csv
import datetime
import io
import json
import re
import zipfile
from html import escape
from types import SimpleNamespace
from uuid import uuid4

from ..models import Artifact, OneKitError, Result, output, parse_json
from .structured_data import bounded, xml_root

MAX_TEXT = 8*1024*1024
_BINARY_CONTROLS = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')


def read_text(file, options):
    try:
        text = file.data.decode(options.get('source_encoding','utf-8-sig'))
    except (UnicodeError,LookupError) as error:
        raise OneKitError('Choose the correct input encoding, or convert the file to UTF-8 first.') from error
    if _BINARY_CONTROLS.search(text):
        raise OneKitError('Choose a text file without binary control characters.')
    return text.replace('\r\n','\n').replace('\r','\n')


def text_rows(text, delimiter):
    if delimiter=='auto':
        try:
            delimiter = csv.Sniffer().sniff(text[:65536],delimiters=',\t;|').delimiter
        except csv.Error:
            delimiter = '\t'
    if delimiter not in (',','\t',';','|'):
        raise OneKitError('Choose tabs, commas, semicolons, pipes or automatic delimiter detection.')
    rows, columns, cells = [], None, 0
    for row in csv.reader(io.StringIO(text,newline=''),delimiter=delimiter,strict=True):
        if not row:
            continue
        if columns is None:
            columns = len(row)
        if len(row)!=columns:
            raise OneKitError('Use consistent column counts, or choose the correct input delimiter.')
        cells += len(row)
        if len(rows)>=100000 or len(row)>1000 or cells>200000:
            raise OneKitError('Use at most 100,000 rows, 1,000 columns and 200,000 cells.')
        rows.append(row)
    if not rows:
        raise OneKitError('Enter at least one non-empty text row.')
    return rows


def delimited(rows, delimiter):
    stream = io.StringIO(newline='')
    csv.writer(stream,delimiter=delimiter).writerows(rows)
    return stream.getvalue()


def spreadsheet_rows(file, options):
    from openpyxl import load_workbook
    book = load_workbook(io.BytesIO(file.data),read_only=True,data_only=True)
    try:
        if options.get('sheet') and options['sheet'] not in book.sheetnames:
            raise OneKitError('Choose a worksheet name that exists in this workbook.')
        sheet = book[options['sheet']] if options.get('sheet') else book.active
        rows, cells = [], 0
        for row in sheet.iter_rows(values_only=True):
            cells += len(row)
            if len(rows)>=100000 or len(row)>1000 or cells>200000:
                raise OneKitError('Use at most 100,000 rows, 1,000 columns and 200,000 cells.')
            rows.append(['' if value is None else value.isoformat() if isinstance(value,(datetime.date,datetime.datetime)) else str(value) for value in row])
        return rows
    finally:
        book.close()


def slide_text(file):
    from .epub_pdf import local_path
    with zipfile.ZipFile(io.BytesIO(file.data)) as archive:
        presentation = xml_root(archive.read('ppt/presentation.xml'))
        relations = xml_root(archive.read('ppt/_rels/presentation.xml.rels'))
        targets = {node.get('Id'):node for node in relations}
        slides = []
        characters = 0
        for node in presentation.iter():
            if node.tag.rsplit('}',1)[-1]!='sldId':
                continue
            ref = node.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
            relation = targets.get(ref)
            if relation is None or relation.get('TargetMode')=='External':
                raise OneKitError('Choose a PowerPoint with local slide relationships.')
            target = relation.get('Target','')
            path = local_path('',target[1:]) if target.startswith('/') and not target.startswith('//') else local_path('ppt',target)
            root = xml_root(archive.read(path))
            paragraphs = []
            for paragraph in root.iter():
                if paragraph.tag=='{http://schemas.openxmlformats.org/drawingml/2006/main}p':
                    paragraphs.append(''.join((part.text or '') if part.tag.endswith('}t') else '\n' if part.tag.endswith('}br') else '' for part in paragraph.iter()))
            slides.append('\n'.join(paragraphs))
            characters += len(slides[-1])
            if len(slides)>256 or characters>MAX_TEXT:
                raise OneKitError('Export at most 256 slides and 8 MiB of text.')
    return '\n\n'.join(slides),len(slides)


def validate_structured(text, kind):
    if kind=='json':
        try:
            value = bounded(parse_json(text))
        except json.JSONDecodeError as error:
            raise OneKitError('Enter valid JSON in the TXT file; arbitrary prose cannot be inferred as structured JSON.') from error
        return json.dumps(value,ensure_ascii=False,allow_nan=False,indent=2)
    if kind=='xml':
        xml_root(text)
        return re.sub(r'(<\?xml\b[^?]*\bencoding\s*=\s*)([\'"])[^\'"]*\2',r'\1"utf-8"',text,count=1,flags=re.IGNORECASE)
    from .structured_data import run as structured
    try:
        structured(SimpleNamespace(id='yaml-to-json'),[Artifact('input.yaml',text.encode())],{})
    except OneKitError:
        raise
    except Exception as error:
        raise OneKitError('Enter valid YAML with consistent indentation and JSON-compatible values.') from error
    return text


def run(tool, files, options):
    file = files[0];source,target = tool.id.split('-to-')
    details, warnings = {}, []
    if target=='txt':
        if source in ('md','html','docx'):
            from .plain_text import docx_text, html_text
            if source=='docx':
                text = docx_text(file.data)
            elif source=='html':
                text = html_text(file.text())
            else:
                from markdown_it import MarkdownIt
                parser = MarkdownIt('commonmark',{'html':True,'maxNesting':32}).enable(['table','strikethrough'])
                blocks = {}
                marker = 'ONEKIT_LITERAL_'+uuid4().hex+'_'
                def literal(tokens,index,render_options,env):
                    key = marker+str(index)
                    blocks[key] = tokens[index].content.rstrip('\n')
                    return '<p>'+key+'</p>\n'
                parser.renderer.rules['fence'] = literal
                parser.renderer.rules['code_block'] = literal
                text = html_text(parser.render(file.text()))
                for key,value in blocks.items():
                    text = text.replace(key,value)
        elif source=='pptx':
            text,details['slides'] = slide_text(file)
        elif source=='epub':
            from .epub_pdf import epub_html
            from .plain_text import html_text
            html,details['chapters'] = epub_html(file)
            text = html_text(html)
        elif source in ('csv','xlsx'):
            rows = spreadsheet_rows(file,options) if source=='xlsx' else text_rows(file.text(),options.get('delimiter',','))
            text = delimited(rows,'\t');details['rows'] = len(rows);details['delimiter'] = 'tab'
        else:
            text = validate_structured(file.text(),source)
        if not text.strip():
            warnings.append('No readable text was found. Images and scanned content require OCR.')
        data = text.encode('utf-8')
        if len(data)>MAX_TEXT:
            raise OneKitError('Extracted text exceeds 8 MiB. Split the source document.')
        return Result([output(file,'txt',data)],warnings,{**details,'characters':len(text),'outputEncoding':'UTF-8'})
    text = read_text(file,options)
    if target in ('json','xml','yaml'):
        return Result([output(file,target,validate_structured(text,target))])
    if target in ('csv','xlsx'):
        rows = text_rows(text,options.get('delimiter','\t'))
        if target=='csv':
            return Result([output(file,'csv',delimited(rows,','))],['CSV values are preserved. Import untrusted values as text in spreadsheet applications.'],{'rows':len(rows)})
        from openpyxl import Workbook
        if any(len(value)>32767 for row in rows for value in row):
            raise OneKitError('Excel cells must contain at most 32,767 characters; shorten the text fields.')
        book = Workbook();sheet = book.active;sheet.title = 'Text'
        for row in rows:
            sheet.append(row)
        for row in sheet:
            for cell in row:
                cell.data_type = 's'
        stream = io.BytesIO();book.save(stream)
        return Result([output(file,'xlsx',stream.getvalue())],details={'rows':len(rows)})
    if target=='md':
        longest = max((len(match.group()) for match in re.finditer(r'`+',text)),default=0)
        fence = '`'*max(3,longest+1)
        return Result([output(file,'md',fence+'\n'+text+('' if text.endswith('\n') else '\n')+fence+'\n')])
    if target=='docx':
        from .converters import text_document
        return Result([text_document(text,file,'docx')])
    if target=='html':
        html = '<!doctype html><html><head><meta charset="utf-8"><title>'+escape(file.stem)+'</title></head><body><pre style="white-space:pre-wrap;overflow-wrap:anywhere">'+escape(text)+'</pre></body></html>'
        return Result([output(file,'html',html)])
    if target=='pdf':
        from .portable_documents import convert
        html = ''.join('<p>'+escape(line or '\u00a0')+'</p>' for line in text.split('\n'))
        return convert(Artifact(file.stem+'.html',html.encode()),'html','pdf')
    raise OneKitError('Choose a supported plain-text conversion.')
